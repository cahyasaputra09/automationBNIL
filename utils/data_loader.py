from itertools import zip_longest
import openpyxl
import os

def load_excel_data(*sheet_names):
    file_path = os.path.join("assets", "testDataBNIL.xlsx")
    workbook = openpyxl.load_workbook(file_path)

    data = []

    # Kalau tidak ada sheet_names, ambil semua sheet
    if not sheet_names:
        sheet_names = workbook.sheetnames

    for sheet_name in sheet_names:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' does not exist in the workbook.")

        sheet = workbook[sheet_name]
        headers = [str(cell.value).strip() if cell.value else "" for cell in sheet[1]]

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row is None or all(cell is None for cell in row):
                continue

            row_dict = {}
            for i in range(len(headers)):
                header = headers[i]
                value = row[i] if i < len(row) else ""
                row_dict[header] = str(value).strip() if value is not None else ""
            data.append(row_dict)

    return data


def load_excel_data_merge(*sheet_names):
    file_path = os.path.join("assets", "testDataBNIL.xlsx")
    workbook = openpyxl.load_workbook(file_path)

    # Ambil semua sheet yang dipilih
    sheets = []
    headers_list = []

    for sheet_name in sheet_names:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' does not exist in the workbook.")
        sheet = workbook[sheet_name]
        sheets.append(sheet.iter_rows(min_row=2, values_only=True))
        headers_list.append([str(cell.value).strip() if cell.value else "" for cell in sheet[1]])

    data = []
    for rows in zip_longest(*sheets, fillvalue=None):
        row_dict = {}
        for sheet_idx, row in enumerate(rows):
            headers = headers_list[sheet_idx]
            if row is None:
                continue
            for i, header in enumerate(headers):
                value = row[i] if i < len(row) and row[i] is not None else ""
                row_dict[header] = str(value).strip()
        if any(value != "" for value in row_dict.values()):
            data.append(row_dict)

    return data

